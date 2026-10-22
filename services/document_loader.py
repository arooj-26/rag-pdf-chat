"""Bounded local loaders; non-paginated sources never acquire invented pages."""
import re
import zipfile
from pathlib import Path
from typing import Protocol
from pypdf import PdfReader
from config import get_settings
from services.pdf_loader import ExtractedPDF, ExtractedPage, PDFExtractionError


class DocumentLoader(Protocol):
    def __call__(self, path: Path) -> ExtractedPDF: ...


def _pdf(path: Path) -> ExtractedPDF:
    settings = get_settings()
    with path.open("rb") as stream:
        if stream.read(5) != b"%PDF-":
            raise PDFExtractionError("The content is not a PDF.")
    reader = PdfReader(path)
    if reader.is_encrypted:
        raise PDFExtractionError("Password-protected PDFs are unsupported. Upload an unlocked copy.")
    if len(reader.pages) > settings.max_document_pages:
        raise PDFExtractionError("Document exceeds the page limit.")
    pages, total, renderer = [], 0, None
    try:
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if len(text.strip()) < 30 and settings.ocr_enabled:
                import pytesseract
                import pypdfium2
                if renderer is None:
                    renderer = pypdfium2.PdfDocument(path)
                pdf_page = renderer[i]
                width, height = pdf_page.get_size()
                if width * height * 4 > 30_000_000:
                    pdf_page.close()
                    raise PDFExtractionError("Page is too large for local OCR.")
                bitmap = pdf_page.render(scale=2)
                try:
                    text = pytesseract.image_to_string(bitmap.to_pil(), lang=settings.ocr_language, timeout=30)
                except (pytesseract.TesseractNotFoundError, RuntimeError) as exc:
                    raise PDFExtractionError("Local OCR is unavailable or timed out. Install Tesseract or disable OCR.") from exc
                finally:
                    bitmap.close()
                    pdf_page.close()
            total += len(text)
            if total > settings.max_extracted_characters:
                raise PDFExtractionError("Extracted text exceeds the processing limit.")
            pages.append(ExtractedPage(i + 1, text))
    finally:
        if renderer is not None:
            renderer.close()
    return ExtractedPDF(pages, len(pages))


def _docx(path: Path) -> ExtractedPDF:
    from docx import Document
    with zipfile.ZipFile(path) as archive:
        if "word/document.xml" not in archive.namelist():
            raise PDFExtractionError("The content is not a Word document.")
        if sum(i.file_size for i in archive.infolist()) > 50_000_000:
            raise PDFExtractionError("Expanded Word document exceeds the limit.")
    document = Document(path)
    sections, lines, heading = [], [], "Introduction"
    for paragraph in document.paragraphs:
        if paragraph.style.name.startswith("Heading"):
            if lines:
                sections.append(ExtractedPage(None, "\n".join(lines), heading))
            heading, lines = paragraph.text, [paragraph.text]
        else:
            lines.append(paragraph.text)
    for table in document.tables:
        lines.extend(" | ".join(c.text for c in row.cells) for row in table.rows)
    sections.append(ExtractedPage(None, "\n".join(lines), heading))
    return ExtractedPDF(sections, 0, "docx")


def _text(path: Path) -> ExtractedPDF:
    raw = path.read_bytes()
    text = raw.decode("utf-8-sig")
    if any(ord(c) < 32 and c not in "\n\r\t" for c in text):
        raise PDFExtractionError("Text files must contain UTF-8 text without binary content.")
    if raw.startswith((b"%PDF", b"PK\x03\x04", b"MZ")):
        raise PDFExtractionError("File content does not match its text extension.")
    sections, heading, lines = [], "Introduction", []
    for line in text.splitlines():
        if path.suffix.lower() == ".md" and re.match(r"^#{1,6}\s", line):
            if lines:
                sections.append(ExtractedPage(None, "\n".join(lines), heading))
            heading, lines = line.lstrip("# "), [line]
        else:
            lines.append(line)
    sections.append(ExtractedPage(None, "\n".join(lines), heading))
    return ExtractedPDF(sections, 0, path.suffix.lower().lstrip("."))


LOADERS: dict[str, DocumentLoader] = {".pdf": _pdf, ".docx": _docx, ".txt": _text, ".md": _text}


def load_document(path: Path, suffix: str | None = None) -> ExtractedPDF:
    settings = get_settings()
    if not path.is_file() or path.stat().st_size == 0:
        raise PDFExtractionError("Document is empty or missing.")
    if path.stat().st_size > settings.max_upload_size_bytes:
        raise PDFExtractionError("Document exceeds the upload limit.")
    loader = LOADERS.get((suffix or path.suffix).lower())
    if loader is None:
        raise PDFExtractionError("Supported formats: PDF, DOCX, TXT and Markdown.")
    try:
        result = loader(path)
    except PDFExtractionError:
        raise
    except Exception as exc:
        raise PDFExtractionError("Document is corrupt or has invalid content for its format.") from exc
    if result.is_empty():
        raise PDFExtractionError("No extractable text. For scanned PDFs enable local OCR.")
    if len(result.full_text) > settings.max_extracted_characters:
        raise PDFExtractionError("Extracted text exceeds the processing limit.")
    return result
