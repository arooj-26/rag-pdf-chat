import pytest
from docx import Document
from pypdf import PdfWriter
from services.document_loader import load_document
from services.pdf_loader import PDFExtractionError
from services.chunking_service import chunk_document


@pytest.mark.parametrize("suffix", [".txt", ".md"])
def test_text_metadata(tmp_path, suffix):
    path = tmp_path / ("notes" + suffix)
    path.write_text("# Biology\nMitochondria generate ATP.", encoding="utf-8")
    result = load_document(path)
    chunks = chunk_document(result, "doc", path.name, "alice")
    assert chunks[0].metadata["owner_id"] == "alice"
    assert chunks[0].metadata["page"] is None
    assert chunks[0].metadata["chunk_id"] == "doc:0"


def test_docx(tmp_path):
    path = tmp_path / "notes.docx"
    doc = Document()
    doc.add_heading("Physics", 1)
    doc.add_paragraph("Energy is conserved.")
    doc.save(path)
    assert "Energy is conserved" in load_document(path).full_text


@pytest.mark.parametrize("name,content", [("empty.txt", b""), ("bad.pdf", b"garbage"), ("bad.docx", b"PKbroken"), ("bad.txt", b"\x00\xff"), ("bad.md", b"%PDF-hello"), ("bad.exe", b"executable")])
def test_invalid(tmp_path, name, content):
    path = tmp_path / name
    path.write_bytes(content)
    with pytest.raises(PDFExtractionError):
        load_document(path)


def test_encrypted_pdf(tmp_path):
    path = tmp_path / "private.pdf"
    writer = PdfWriter()
    writer.add_blank_page(100, 100)
    writer.encrypt("password")
    writer.write(path)
    with pytest.raises(PDFExtractionError, match="Password-protected"):
        load_document(path)
